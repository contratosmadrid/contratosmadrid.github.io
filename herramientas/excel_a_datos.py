#!/usr/bin/env python3
"""
Convierte el Excel maestro (hoja «Datos») en los datos que lee el mapa web.

Entrada : datos/*.xlsx  (debe haber exactamente un Excel)
Salida  : web/data/contratos.js   -> lo carga index.html (funciona también abriendo el archivo en local)
          web/data/contratos.json -> mismo contenido, para reutilizar
          web/data/informe.md     -> resumen de la conversión y de los cambios respecto a la versión anterior

Si encuentra un error estructural (columnas que faltan, ID REL repetidos o vacíos,
campos obligatorios vacíos, distritos desconocidos, fechas ilegibles) termina con
código 1 y NO se publica nada: la web anterior sigue en línea.
"""
import datetime as dt
import glob
import json
import os
import re
import sys
from zoneinfo import ZoneInfo

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "web", "data")
HEREDADAS = os.path.join(RAIZ, "datos", "particularidades_heredadas.json")

DISTRITOS = ['Centro', 'Arganzuela', 'Retiro', 'Salamanca', 'Chamartín', 'Tetuán', 'Chamberí',
             'Fuencarral-El Pardo', 'Moncloa-Aravaca', 'Latina', 'Carabanchel', 'Usera',
             'Puente de Vallecas', 'Moratalaz', 'Ciudad Lineal', 'Hortaleza', 'Villaverde',
             'Villa de Vallecas', 'Vicálvaro', 'San Blas-Canillejas', 'Barajas']
AMBITO = 'Ámbito municipal'

# campo de la web  ->  columna del Excel (hoja «Datos»)
COLUMNAS = {
    'distrito': 'Distrito', 'servicio': 'Servicio', 'area': 'Area de gobierno',
    'direccion': 'Direccion general', 'unidad': 'Subdireccion / unidad', 'contrato': 'Contrato',
    'expediente': 'Expediente', 'lote': 'Lote', 'empresa': 'Empresa / UTE', 'prestacion': 'Prestacion',
    'cobertura': 'Cobertura', 'inicio': 'Inicio registrado (ver criterio)',
    'fin': 'Fin registrado (ver criterio)', 'responsable': 'Responsable DG',
    'gestor': 'Gestor del contrato', 'contacto': 'Contacto publicado',
    'excepciones': 'Excepciones / limites', 'pendientes': 'Estado / pendientes', 'fuente': 'Fuente',
    'relid': 'ID relacion (PK propuesta)', 'loteid': 'ID lote (FK propuesta)',
    'vigencia': 'ESTADO DE VIGENCIA · 16/10/2026', 'vinculo': 'CRITERIO VÍNCULO CORPORATIVO',
    'grupos': 'CLAVES DE GRUPO',
}
OBLIGATORIOS = ['relid', 'distrito', 'expediente', 'contrato', 'empresa', 'servicio']
ORDEN = ['distrito', 'servicio', 'area', 'direccion', 'unidad', 'contrato', 'expediente', 'lote',
         'empresa', 'prestacion', 'cobertura', 'inicio', 'fin', 'responsable', 'gestor', 'contacto',
         'excepciones', 'estado', 'fuente', 'relid', 'loteid', 'grupos', 'vigencia', 'tipo', 'vinculo']


def fecha(v, ctx, errores):
    """Devuelve dd/mm/aaaa o '' si la celda está vacía."""
    if v is None or str(v).strip() == '':
        return ''
    if isinstance(v, dt.datetime):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, dt.date):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # fecha numérica de Excel
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).strftime('%d/%m/%Y')
    s = str(v).strip()
    if re.fullmatch(r'\d{2}/\d{2}/\d{4}', s):
        return s
    errores.append(f'{ctx}: fecha no reconocida «{s}» (use dd/mm/aaaa o formato fecha de Excel)')
    return s


def texto(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def main():
    excels = [f for f in glob.glob(os.path.join(RAIZ, 'datos', '*.xlsx')) if not os.path.basename(f).startswith('~$')]
    if len(excels) != 1:
        print(f'ERROR: en la carpeta datos/ tiene que haber exactamente un Excel (.xlsx). Encontrados: {len(excels)}')
        for e in excels:
            print('   -', os.path.basename(e))
        sys.exit(1)
    origen = excels[0]
    wb = openpyxl.load_workbook(origen, data_only=True, read_only=True)
    if 'Datos' not in wb.sheetnames:
        print('ERROR: el Excel no tiene la hoja «Datos».'); sys.exit(1)
    ws = wb['Datos']
    filas = ws.iter_rows(values_only=True)
    cab = [texto(c) for c in next(filas)]
    faltan = [c for c in COLUMNAS.values() if c not in cab]
    if faltan:
        print('ERROR: faltan columnas en la hoja «Datos»:'); [print('   -', c) for c in faltan]
        print('No cambie el nombre de las columnas sin adaptar herramientas/excel_a_datos.py'); sys.exit(1)
    idx = {k: cab.index(c) for k, c in COLUMNAS.items()}

    heredadas = {}
    if os.path.exists(HEREDADAS):
        heredadas = json.load(open(HEREDADAS, encoding='utf-8')).get('textos', {})

    errores, avisos, registros, vistos = [], [], [], {}
    for n, fila in enumerate(filas, start=2):
        if fila is None or all(v is None or str(v).strip() == '' for v in fila):
            continue
        g = lambda k: fila[idx[k]] if idx[k] < len(fila) else None
        # Filas auxiliares (fórmulas de apoyo) sin datos de contrato: se ignoran
        if all(texto(g(k)) == '' for k in ('relid', 'distrito', 'expediente', 'contrato', 'empresa')):
            continue
        r = {k: texto(g(k)) for k in COLUMNAS if k not in ('inicio', 'fin')}
        ctx = f'Fila {n} ({r["relid"] or "sin ID"})'
        r['inicio'] = fecha(g('inicio'), ctx + ' · Inicio', errores)
        r['fin'] = fecha(g('fin'), ctx + ' · Fin', errores)
        for k in OBLIGATORIOS:
            if not r[k]:
                errores.append(f'{ctx}: campo obligatorio vacío «{COLUMNAS[k]}»')
        if r['relid']:
            if not re.fullmatch(r'REL-\d{4,}', r['relid']):
                errores.append(f'{ctx}: ID de relación con formato incorrecto «{r["relid"]}» (debe ser REL-0001…)')
            if r['relid'] in vistos:
                errores.append(f'{ctx}: ID {r["relid"]} repetido (ya aparece en la fila {vistos[r["relid"]]})')
            vistos[r['relid']] = n
        if r['distrito'] and r['distrito'] not in DISTRITOS and r['distrito'] != AMBITO:
            errores.append(f'{ctx}: distrito desconocido «{r["distrito"]}»')
        # Igual que la web v2: «estado» = vigencia + pendientes
        r['estado'] = r['vigencia'] + ' · ' + r['pendientes']
        if not r['excepciones'] and r['relid'] in heredadas:
            r['excepciones'] = heredadas[r['relid']]
        r['tipo'] = ''
        if not re.search(r'https?://', r['fuente']):
            avisos.append(f'{ctx}: sin enlace de fuente')
        registros.append({k: r[k] for k in ORDEN})

    if errores:
        print(f'ERROR: {len(errores)} problema(s) en el Excel. No se publica nada.\n')
        for e in errores[:200]:
            print(' -', e)
        sys.exit(1)

    # Duplicados lógicos: misma relación expediente+lote+distrito+empresa+prestación
    clave = {}
    for r in registros:
        k = (r['expediente'], r['lote'], r['distrito'], r['empresa'], r['prestacion'])
        if k in clave:
            avisos.append(f'Posible duplicado: {r["relid"]} y {clave[k]} tienen mismo expediente, lote, distrito, empresa y prestación')
        clave[k] = r['relid']

    # Comparación con la versión publicada anterior
    previo_path = os.path.join(SALIDA, 'contratos.json')
    cambios = []
    if os.path.exists(previo_path):
        try:
            previo = {r['relid']: r for r in json.load(open(previo_path, encoding='utf-8'))['records']}
            nuevo = {r['relid']: r for r in registros}
            for k in sorted(set(nuevo) - set(previo)):
                cambios.append(f'➕ Nueva relación {k} · {nuevo[k]["expediente"]} · {nuevo[k]["distrito"]}')
            for k in sorted(set(previo) - set(nuevo)):
                cambios.append(f'➖ Eliminada {k} · {previo[k]["expediente"]} · {previo[k]["distrito"]}')
            for k in sorted(set(previo) & set(nuevo)):
                dif = [c for c in ORDEN if previo[k].get(c) != nuevo[k].get(c)]
                if dif:
                    cambios.append(f'✏️ {k} · {nuevo[k]["expediente"]}: cambia {", ".join(dif)}')
        except Exception as e:  # versión anterior ilegible: no bloquea
            avisos.append(f'No se pudo comparar con la versión anterior: {e}')

    exp = {r['expediente'] for r in registros}
    resumen = {
        'relaciones': len(registros), 'expedientes': len(exp),
        'fichas': len({(r['expediente'], r['lote']) for r in registros}),
        'empresas': len({r['empresa'] for r in registros}),
        'ambito_municipal': sum(r['distrito'] == AMBITO for r in registros),
        'sin_fecha_fin': len({(r['expediente'], r['lote']) for r in registros if not r['fin']}),
    }
    ahora = dt.datetime.now(ZoneInfo('Europe/Madrid'))
    datos = {'generado': ahora.strftime('%d/%m/%Y %H:%M'), 'origen': os.path.basename(origen),
             'resumen': resumen, 'records': registros}

    os.makedirs(SALIDA, exist_ok=True)
    js = json.dumps(datos, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    open(os.path.join(SALIDA, 'contratos.js'), 'w', encoding='utf-8').write(
        '/* Generado automáticamente desde el Excel maestro. No editar a mano. */\nwindow.MADRID_DATA=' + js + ';\n')
    json.dump(datos, open(os.path.join(SALIDA, 'contratos.json'), 'w', encoding='utf-8'), ensure_ascii=False)

    inf = [f'# Informe de conversión · {datos["generado"]}', '', f'Excel: `{datos["origen"]}`', '',
           '| Indicador | Valor |', '|---|---|'] + [f'| {k} | {v} |' for k, v in resumen.items()] + ['']
    inf += ['## Cambios respecto a la versión anterior', ''] + (cambios or ['Sin cambios (o primera publicación).']) + ['']
    inf += [f'## Avisos ({len(avisos)})', ''] + [f'- {a}' for a in avisos[:300]]
    open(os.path.join(SALIDA, 'informe.md'), 'w', encoding='utf-8').write('\n'.join(inf) + '\n')
    if os.environ.get('GITHUB_STEP_SUMMARY'):
        open(os.environ['GITHUB_STEP_SUMMARY'], 'a', encoding='utf-8').write('\n'.join(inf) + '\n')

    print(f'OK · {resumen["expedientes"]} expedientes · {resumen["relaciones"]} relaciones · '
          f'{len(cambios)} cambios · {len(avisos)} avisos')


if __name__ == '__main__':
    main()
